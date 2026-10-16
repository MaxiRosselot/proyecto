"""Genera ventas-ejemplo.xlsx: un libro con el mismo formato que la planilla de Maxi, con datos
inventados, para la prueba E2E de la importacion (verify-planilla.mjs).
  python scripts/local-dev/fixtures/generar-ventas-ejemplo.py
"""
import datetime as dt
from pathlib import Path
import openpyxl
from openpyxl.styles import PatternFill

GRIS = PatternFill('solid', fgColor='FFEFEFEF')
libro = openpyxl.Workbook()

# Como "VENTAS ENERO" 2025: REALIZADO y DIA adelante, W de Webpay.
h = libro.active
h.title = 'VENTAS ENERO'
h.append(['', '', '', '', '', '', '', '', 'MONTO NETO'])
h.append(['REALIZADO', 'DÍA', 'NOMBRE', 'FECHA', 'HORA INICIO', 'HORA FIN', 'B/F/W', 'PAGADO', 'REPISAS', 'RETIRO', 'SOPORTE', 'ESCOMBROS', 'NETO TOTAL', 'MONTO BRUTO'])
for fila in [
    [False, 'MARTES', 'TOMAS PRUEBA', dt.datetime(2025, 1, 14), dt.time(14), dt.time(16), 'W', 'SI', 170000, 0, 0, 0, 170000, 202300],
    [False, 'SÁBADO', 'CAMILA EJEMPLO', dt.datetime(2025, 1, 18), dt.time(14), dt.time(16), 'B', 'SI', 140000, 30000, 0, 30000, 200000, 238000],
]:
    h.append(fila)
    h.cell(h.max_row, 3).fill = GRIS

# Como "VENTAS 2026": nota en DIA, casilla en B, DESPACHO y RACKS, y un resumen a la derecha.
h = libro.create_sheet('VENTAS 2026')
h.append(['', False, '', '', '', '', '', '', '', 'MONTO NETO'])
h.append(['DÍA', False, 'NOMBRE', 'FECHA', 'HORA INICIO', 'HORA FIN', '', 'B/F/W', 'PAGADO', 'REPISAS', 'RETIRO', 'SOPORTE', 'ESCOMBROS', 'DESPACHO', 'RACKS', 'NETO TOTAL', 'MONTO BRUTO', '', 'CLIENTES', 'TOTAL', 'PAGADO'])
filas = [
    (['CAJAS PENDIENTES', False, 'IGNACIA DEMO', dt.datetime(2026, 9, 16), dt.time(13), dt.time(16), '', 'B*', 'SI', 300000, 40000, 0, 0, 0, 0, 340000, 404600], True),
    (['', False, 'MARCO PRUEBA', dt.datetime(2026, 9, 22), dt.time(10), dt.time(13), '', 'M', 'NO', 190000, 0, 0, 0, 0, 0, 190000, 226100], True),
    (['VALPARAISO', False, 'LUCIA FICTICIA', None, None, None, '', 'b/m', 'SI', 140000, 0, 0, 0, 60000, 0, 200000, 238000], True),
    (['', False, 'PEDRO FUTURO', dt.datetime(2027, 3, 10), dt.time(10), dt.time(13), '', '', '', 0, 0, 0, 0, 0, 0, 0, 0], False),
]
for valores, pintada in filas:
    h.append(valores)
    if pintada: h.cell(h.max_row, 3).fill = GRIS

# Una hoja del banco: no es de ventas y se ignora.
h = libro.create_sheet('ENERO 2025')
h.append(['', 1050667])
h.append(['FECHA', 'MOVIMIENTO', 'TIPO', 'MONTO', 'SALDO'])
h.append([dt.datetime(2025, 1, 2), 'VENTA WEBPAY', 'ABONO', 99585, 2027953])

destino = Path(__file__).with_name('ventas-ejemplo.xlsx')
libro.save(destino)
print(destino)
